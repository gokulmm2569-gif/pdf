import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Sr No",
    "Doc No",
    "Date",
    "Dest. City",
    "Mode",
    "I",
    "PC",
    "Weight",
    "Amount"
]

def create_excel_workbook(records: list[dict], output_path: str) -> str:
    """
    Creates a professionally formatted .xlsx file strictly complying with:
    - Sheet name: Invoice_Data
    - First row: Sr No, Doc No, Date, Dest. City, Mode, I, PC, Weight, Amount
    - Header row bold with styling
    - Columns properly sized
    - Weight as numeric value where possible
    - Amount as numeric value where possible
    - Preserve leading zeros on Doc No (explicit string type)
    - Thin borders and clean typography
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice_Data"
    
    # Enable grid lines visibility
    ws.views.sheetView[0].showGridLines = True
    
    # Header styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    
    thin_border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )
    
    # Write header row
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = cell_border
        
        # Header alignments
        if col_name in ["Sr No", "Date", "Mode", "I", "PC"]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        elif col_name in ["Weight", "Amount"]:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            
    ws.row_dimensions[1].height = 24
    
    # Write data records
    for row_idx, rec in enumerate(records, 2):
        ws.row_dimensions[row_idx].height = 20
        
        # 1. Sr No
        sr_val = rec.get("Sr No", "")
        sr_cell = ws.cell(row=row_idx, column=1)
        if str(sr_val).strip().isdigit():
            sr_cell.value = int(sr_val)
        else:
            sr_cell.value = str(sr_val).strip()
        sr_cell.alignment = Alignment(horizontal="center", vertical="center")
        sr_cell.font = data_font
        sr_cell.border = cell_border
        
        # 2. Doc No (Store strictly as string to preserve leading zeros)
        doc_val = str(rec.get("Doc No", "")).strip()
        doc_cell = ws.cell(row=row_idx, column=2, value=doc_val)
        doc_cell.number_format = "@" # Text format
        doc_cell.alignment = Alignment(horizontal="left", vertical="center")
        doc_cell.font = data_font
        doc_cell.border = cell_border
        
        # 3. Date
        date_val = str(rec.get("Date", "")).strip()
        date_cell = ws.cell(row=row_idx, column=3, value=date_val)
        date_cell.alignment = Alignment(horizontal="center", vertical="center")
        date_cell.font = data_font
        date_cell.border = cell_border
        
        # 4. Dest. City
        dest_val = str(rec.get("Dest. City", "")).strip()
        dest_cell = ws.cell(row=row_idx, column=4, value=dest_val)
        dest_cell.alignment = Alignment(horizontal="left", vertical="center")
        dest_cell.font = data_font
        dest_cell.border = cell_border
        
        # 5. Mode
        mode_val = str(rec.get("Mode", "")).strip()
        mode_cell = ws.cell(row=row_idx, column=5, value=mode_val)
        mode_cell.alignment = Alignment(horizontal="center", vertical="center")
        mode_cell.font = data_font
        mode_cell.border = cell_border
        
        # 6. I
        i_val = str(rec.get("I", "")).strip()
        i_cell = ws.cell(row=row_idx, column=6, value=i_val)
        i_cell.alignment = Alignment(horizontal="center", vertical="center")
        i_cell.font = data_font
        i_cell.border = cell_border
        
        # 7. PC
        pc_val = rec.get("PC", "")
        pc_cell = ws.cell(row=row_idx, column=7)
        if str(pc_val).strip().isdigit():
            pc_cell.value = int(pc_val)
        else:
            pc_cell.value = str(pc_val).strip()
        pc_cell.alignment = Alignment(horizontal="center", vertical="center")
        pc_cell.font = data_font
        pc_cell.border = cell_border
        
        # 8. Weight (Numeric with 3 decimals if convertible)
        w_val = str(rec.get("Weight", "")).replace(",", "").strip()
        w_cell = ws.cell(row=row_idx, column=8)
        try:
            w_cell.value = float(w_val)
            w_cell.number_format = "0.000"
        except (ValueError, TypeError):
            w_cell.value = w_val
        w_cell.alignment = Alignment(horizontal="right", vertical="center")
        w_cell.font = data_font
        w_cell.border = cell_border
        
        # 9. Amount (Numeric with 2 decimals if convertible)
        amt_val = str(rec.get("Amount", "")).replace(",", "").strip()
        amt_cell = ws.cell(row=row_idx, column=9)
        try:
            amt_cell.value = float(amt_val)
            amt_cell.number_format = "0.00"
        except (ValueError, TypeError):
            amt_cell.value = amt_val
        amt_cell.alignment = Alignment(horizontal="right", vertical="center")
        amt_cell.font = data_font
        amt_cell.border = cell_border
        
    # Auto-fit column widths with safety margin
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    wb.save(output_path)
    return output_path
